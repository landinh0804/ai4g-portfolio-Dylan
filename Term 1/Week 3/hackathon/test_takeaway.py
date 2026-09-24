"""Tests for the take-away sheet.

The sheet is the only output that leaves the screen, so the things worth asserting
are the ones that would be wrong on paper in somebody's hand: that the Dutch column
is there, that a quote is reproduced exactly as it appears in the contract, that a
gated run does not read as a summary, and that an unverified rule keeps its caveat.
"""

from __future__ import annotations

from datetime import datetime, timezone

import pytest

from contract_trap_finder.models import (
    ContactOption,
    ExitEvent,
    ExitScenario,
    Finding,
    Report,
    SafeQuestion,
    SourceQuote,
)
from contract_trap_finder import takeaway

DUTCH_CLAUSE = "De uitzendkracht is een bemiddelingsvergoeding verschuldigd van EUR 350,00."


def question(**overrides) -> SafeQuestion:
    data = {
        "question_in_user_language": "Czy muszę zapłacić za znalezienie tej pracy?",
        "question_in_employer_language": "Moet ik betalen voor het vinden van dit werk?",
        "why_this_question": "A fee charged to you for the placement is void under Waadi article 9.",
        "what_a_refusal_means": "If this is not answered in writing, keep the question and the date.",
    }
    data.update(overrides)
    return SafeQuestion(**data)


def finding(**overrides) -> Finding:
    data = {
        "id": "NL-WAADI-09-FEE",
        "category": "ILLEGAL",
        "severity": "high",
        "title": "Placony jest od ciebie fee za pracę",
        "what_it_means": "You are being charged for being placed in this job.",
        "quotes": [SourceQuote(doc_id="01_uitzendovereenkomst.txt", quote=DUTCH_CLAUSE)],
        "origin": "deterministic",
        "rule_id": "NL-WAADI-09-FEE",
        "rule_verified": True,
    }
    data.update(overrides)
    return Finding(**data)


def report(**overrides) -> Report:
    data = {
        "generated_at": datetime(2026, 9, 19, tzinfo=timezone.utc),
        "user_language": "pl",
        "findings": [finding()],
        "safe_questions": [question()],
        "contacts": [
            ContactOption(
                name="FairWork",
                kind="confidential_advice",
                what_they_do="Advice for migrant workers, in confidence.",
                will_employer_find_out="No, not unless you ask them to act.",
                typical_response_time="A few days",
                contact="+31 20 000 0000 (check the website, this number is unverified)",
                url="https://www.fairwork.nu",
            )
        ],
        "overall_confidence": 0.82,
    }
    data.update(overrides)
    return Report(**data)


def text_of(document) -> str:
    """Everything the reader can see, including inside tables."""
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_the_sheet_is_bilingual() -> None:
    """The worker's language and Dutch both appear, for the same question.

    This is the whole point of the sheet: somebody who speaks no Dutch can put it in
    front of a Dutch speaker and point. A sheet in one language is a printout.
    """
    content = text_of(takeaway.build_document(report()))
    assert "Czy muszę zapłacić za znalezienie tej pracy?" in content
    assert "Moet ik betalen voor het vinden van dit werk?" in content


def test_quotes_are_reproduced_exactly() -> None:
    """The quote is the thing that can be checked against the contract.

    safety.verify_quote has already confirmed this string appears in the source
    document. If the sheet reformats it, the worker points at a line that is not in
    their contract.
    """
    content = text_of(takeaway.build_document(report()))
    assert DUTCH_CLAUSE in content


def test_static_labels_appear_in_the_workers_language() -> None:
    content = text_of(takeaway.build_document(report(user_language="pl")))
    assert takeaway.LABELS["pl"]["questions"] in content
    assert takeaway.LABELS["pl"]["title"] in content


def test_dutch_title_is_added_when_the_worker_reads_another_language() -> None:
    content = text_of(takeaway.build_document(report(user_language="pl")))
    assert takeaway.LABELS["nl"]["title"] in content


def test_a_dutch_reader_does_not_get_the_title_twice() -> None:
    content = text_of(takeaway.build_document(report(user_language="nl")))
    assert content.count(takeaway.LABELS["nl"]["title"]) == 1


def test_unchecked_label_translations_say_so() -> None:
    """Polish labels have not been read by a native speaker, and the sheet admits it."""
    content = text_of(takeaway.build_document(report(user_language="pl")))
    assert takeaway.LABELS["en"]["labels_unchecked"] in content or takeaway.LABELS["pl"].get("labels_unchecked", "") in content

    checked = text_of(takeaway.build_document(report(user_language="en")))
    assert takeaway.LABELS["en"]["labels_unchecked"] not in checked


def test_a_gated_run_leads_with_the_refusal() -> None:
    """A refusal at the bottom of the page is not a refusal."""
    gated = report(
        gated=True,
        gate_reason="The two documents state different hourly wages.",
        conflicts=["Hourly wage: EUR 15,10 against EUR 12,85"],
        user_language="en",
    )
    document = takeaway.build_document(gated)
    content = text_of(document)

    assert takeaway.LABELS["en"]["gated"] in content
    assert "The two documents state different hourly wages." in content
    assert "Hourly wage: EUR 15,10 against EUR 12,85" in content

    banner = content.index(takeaway.LABELS["en"]["gated"])
    findings_heading = content.index(takeaway.LABELS["en"]["findings"])
    assert banner < findings_heading, "the refusal must come before the findings"


def test_deterministic_findings_survive_the_gate() -> None:
    """A rule check that fired is still a fact, which is what pipeline.py says."""
    content = text_of(takeaway.build_document(report(gated=True, gate_reason="Documents conflict.")))
    assert DUTCH_CLAUSE in content


def test_unverified_rules_keep_their_caveat() -> None:
    unverified = report(findings=[finding(rule_verified=False)], user_language="en")
    content = text_of(takeaway.build_document(unverified))
    assert takeaway.LABELS["en"]["unverified"] in content

    verified = report(findings=[finding(rule_verified=True)], user_language="en")
    assert takeaway.LABELS["en"]["unverified"] not in text_of(takeaway.build_document(verified))


def test_the_timeline_is_included_when_there_is_one() -> None:
    with_timeline = report(
        user_language="en",
        exit_scenario=ExitScenario(
            trigger_description="If your assignment ends",
            events=[
                ExitEvent(day_offset=0, label="Income stops", description="No further hours are offered."),
                ExitEvent(day_offset=2, label="Housing ends", description="You are asked to leave the room."),
            ],
        ),
    )
    content = text_of(takeaway.build_document(with_timeline))
    assert "If your assignment ends" in content
    assert "Housing ends" in content


def test_an_empty_report_still_produces_a_sheet() -> None:
    """A run that found nothing still has questions worth asking and contacts."""
    document = takeaway.build_document(Report(generated_at=datetime.now(timezone.utc), user_language="en"))
    assert takeaway.LABELS["en"]["title"] in text_of(document)


def test_no_model_call_is_needed(monkeypatch: pytest.MonkeyPatch) -> None:
    """The sheet must be buildable with no key, no network and no Ollama.

    Asserted rather than assumed, because a later change that reaches for a
    translation here would put an unchecked sentence on a page built to be handed to
    an employer.
    """
    import contract_trap_finder.llm as llm

    def explode(*args, **kwargs):
        raise AssertionError("build_document must not construct an LLM client")

    monkeypatch.setattr(llm, "build_client", explode)
    takeaway.build_document(report())


def test_write_takeaway_produces_a_readable_file(tmp_path) -> None:
    from docx import Document as OpenDocument

    path = takeaway.write_takeaway(report(), tmp_path / "sheet.docx")
    assert path.exists() and path.stat().st_size > 0

    reopened = OpenDocument(str(path))
    assert DUTCH_CLAUSE in text_of(reopened)


def test_the_contact_route_and_its_caveat_are_printed() -> None:
    """A contact nobody can reach is not a contact.

    rules/contacts.json marks unverified details in the string itself, and that
    caveat has to survive onto the page - the sheet is the copy that gets kept, so it
    is the copy most likely to be acted on weeks later.
    """
    content = text_of(takeaway.build_document(report()))
    assert "+31 20 000 0000 (check the website, this number is unverified)" in content
    assert "https://www.fairwork.nu" in content


def test_default_filename_names_the_language_and_date() -> None:
    name = takeaway.default_filename(report(user_language="pl"))
    assert name.startswith("questions-to-ask-pl-2026-09-19")
    assert name.endswith(".docx")


def test_every_supported_language_has_labels() -> None:
    """A language offered in the UI must not fall back to English silently."""
    from contract_trap_finder import config

    missing = [code for code in config.SUPPORTED_LANGUAGES if code not in takeaway.LABELS]
    assert not missing, f"the take-away sheet has no labels for {missing}"


def test_label_blocks_all_carry_the_same_keys() -> None:
    expected = set(takeaway.LABELS["en"])
    for code, block in takeaway.LABELS.items():
        assert set(block) == expected, f"{code} labels differ from the English set: {set(block) ^ expected}"


def test_the_sheet_tells_the_reader_to_talk_to_a_person() -> None:
    """On paper a list of quoted clauses under a heading reads as a verdict.

    The sheet is the output that gets printed and handed over, so the line saying a
    computer did the reading has to be on it, in the reader's own language.
    """
    document = takeaway.build_document(report())
    body = text_of(document)
    assert takeaway.LABELS["pl"]["talk_to_human"] in body


def test_the_talk_to_a_person_line_is_above_the_findings() -> None:
    """Under the findings it would be a footnote. Order is the whole point."""
    body = text_of(takeaway.build_document(report()))
    assert body.index(takeaway.LABELS["pl"]["talk_to_human"]) < body.index(
        report().findings[0].title
    )
